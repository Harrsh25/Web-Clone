from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

TEMPLATE_PATH = "/root/.claude/uploads/b5947347-19f1-59d7-8db3-07624bf3b17b/74b734a8-Inventia_Template_V1.0_2.pptx"
OUTPUT_PATH = "/home/user/Web-Clone/WFM_Platform_v2.pptx"

# Colors
C_BLUE = RGBColor(0x00, 0x70, 0xC0)
C_TEAL = RGBColor(0x12, 0xA8, 0xBC)
C_BODY = RGBColor(0x1C, 0x2A, 0x38)
C_MUTED = RGBColor(0x5E, 0x6E, 0x7C)
C_LIGHT = RGBColor(0xF4, 0xF7, 0xFA)
C_DIVIDER = RGBColor(0xE2, 0xE9, 0xEF)
C_SECNUM = RGBColor(0xE8, 0xF0, 0xF8)
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation(TEMPLATE_PATH)

def del_slide(prs, idx):
    sldIdLst = prs.slides._sldIdLst
    sId = sldIdLst[idx]
    rId = sId.get(qn('r:id'))
    sldIdLst.remove(sId)
    prs.part.drop_rel(rId)

for _ in range(len(prs.slides)):
    del_slide(prs, 0)

def add_slide():
    return prs.slides.add_slide(prs.slide_layouts[1])

def add_rect(slide, l, t, w, h, fill_color, no_border=True):
    shape = slide.shapes.add_shape(1, Emu(l), Emu(t), Emu(w), Emu(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if no_border:
        shape.line.fill.background()
    return shape

def add_oval(slide, l, t, w, h, fill_color, no_border=True):
    shape = slide.shapes.add_shape(9, Emu(l), Emu(t), Emu(w), Emu(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if no_border:
        shape.line.fill.background()
    return shape

def add_textbox(slide, l, t, w, h):
    tb = slide.shapes.add_textbox(Emu(l), Emu(t), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    return tf

def set_para(para, text, size, color, bold=False, italic=False, align=None, space_before=6, space_after=6):
    para.clear()
    run = para.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    para.space_before = Pt(space_before)
    para.space_after = Pt(space_after)
    if align:
        para.alignment = align

def add_text(slide, l, t, w, h, text, size, color, bold=False, italic=False, align=None):
    tf = add_textbox(slide, l, t, w, h)
    para = tf.paragraphs[0]
    set_para(para, text, size, color, bold, italic, align)
    return tf

def top_bar(slide):
    add_rect(slide, 0, 0, 12192000, 50800, C_BLUE)

def divider(slide, l, t, w=10912000):
    add_rect(slide, l, t, w, 12700, C_DIVIDER)

def section_header(slide, label, number):
    add_text(slide, 640080, 100000, 5000000, 200000, label, 11, C_BLUE)
    add_text(slide, 640080, 200000, 1000000, 600000, number, 60, C_SECNUM, bold=True)

def outcome_box(slide, outcome_text):
    add_rect(slide, 7800000, 1200000, 3800000, 3800000, C_LIGHT)
    add_rect(slide, 7800000, 1200000, 50800, 3800000, C_TEAL)
    add_text(slide, 7900000, 1280000, 3600000, 200000, "OUTCOME", 10, C_TEAL, bold=True)
    add_text(slide, 7900000, 1450000, 3600000, 3400000, outcome_text, 13, C_BODY)

def feature_slide(label, number, title, bullets, outcome_text):
    slide = add_slide()
    top_bar(slide)
    section_header(slide, label, number)
    add_text(slide, 640080, 700000, 6800000, 600000, title, 26, C_BODY, bold=True)
    tf = add_textbox(slide, 640080, 1500000, 6800000, 3000000)
    first = True
    for b in bullets:
        if first:
            para = tf.paragraphs[0]
            first = False
        else:
            para = tf.add_paragraph()
        set_para(para, b, 13, C_BODY, space_before=8, space_after=8)
    outcome_box(slide, outcome_text)
    return slide

# ── SLIDE 1: Title ──────────────────────────────────────────────────────────
slide1 = add_slide()
add_rect(slide1, 0, 0, 76200, 6858000, C_BLUE)
add_text(slide1, 609600, 1800000, 9000000, 800000, "Workforce Management Platform", 44, C_BODY, bold=True)
add_text(slide1, 609600, 2700000, 9000000, 400000, "Plan Smarter. Execute Faster. Deliver More.", 20, C_BLUE)
add_rect(slide1, 609600, 3300000, 5000000, 12700, C_BLUE)
add_text(slide1, 609600, 3500000, 3000000, 300000, "Inventia", 14, C_MUTED)
add_text(slide1, 9500000, 6200000, 2000000, 200000, "v1.0 | 2025", 11, C_MUTED, align=PP_ALIGN.RIGHT)

# ── SLIDE 2: Platform Overview ───────────────────────────────────────────────
slide2 = add_slide()
top_bar(slide2)
add_text(slide2, 640080, 200000, 8000000, 400000, "PLATFORM OVERVIEW", 11, C_BLUE)
add_text(slide2, 640080, 900000, 3000000, 700000, "3x Faster", 48, C_BLUE, bold=True)
add_text(slide2, 640080, 1600000, 3000000, 300000, "Planning", 14, C_MUTED)
add_text(slide2, 4500000, 900000, 3000000, 700000, "40%", 48, C_BLUE, bold=True)
add_text(slide2, 4500000, 1600000, 3000000, 300000, "Less Admin", 14, C_MUTED)
add_text(slide2, 8200000, 900000, 3000000, 700000, "Real-Time", 48, C_BLUE, bold=True)
add_text(slide2, 8200000, 1600000, 3000000, 300000, "Visibility", 14, C_MUTED)
divider(slide2, 640080, 2200000)
add_text(slide2, 640080, 2400000, 10912000, 600000,
    "The Workforce Management Platform gives every team member, manager, and executive the clarity and control they need — from individual task planning to portfolio-level visibility.",
    14, C_BODY)

# ── SLIDE 3: The Problem ─────────────────────────────────────────────────────
slide3 = add_slide()
top_bar(slide3)
section_header(slide3, "THE PROBLEM", "03")
add_text(slide3, 640080, 700000, 10000000, 400000, "Why Teams Struggle to Deliver", 28, C_BODY, bold=True)

pain_points = [
    (1400000, "01", "Projects run on spreadsheets and disconnected tools, causing version chaos"),
    (2200000, "02", "Managers lack real-time visibility into who is doing what and when"),
    (3000000, "03", "Approval cycles are manual, slow, and create bottlenecks"),
    (3800000, "04", "No single source of truth for project health, resources, or timelines"),
]
for y, num, desc in pain_points:
    add_text(slide3, 640080, y, 600000, 400000, num, 32, C_BLUE, bold=True)
    add_text(slide3, 1400000, y, 9800000, 400000, desc, 14, C_BODY)

for dy in [2000000, 2800000, 3600000]:
    divider(slide3, 640080, dy)

# ── SLIDE 4: Who It's For ────────────────────────────────────────────────────
slide4 = add_slide()
top_bar(slide4)
section_header(slide4, "WHO IT'S FOR", "04")
add_text(slide4, 640080, 700000, 10000000, 400000, "Built for Every Industry", 28, C_BODY, bold=True)

industries_left = [
    ("IT & Software", "Sprint planning, release tracking, and resource management"),
    ("Manufacturing", "Production scheduling, shift planning, and capacity management"),
    ("Healthcare", "Staff scheduling, compliance tracking, and project governance"),
]
industries_right = [
    ("Construction & EPC", "WBS planning, milestone tracking, and subcontractor management"),
    ("Professional Services", "Client project delivery, utilization tracking, and billing"),
    ("Financial Services", "Regulatory project tracking, audit trails, and risk management"),
]
row_ys = [1500000, 2600000, 3700000]

for i, (name, desc) in enumerate(industries_left):
    y = row_ys[i]
    add_text(slide4, 640080, y, 5000000, 300000, name, 16, C_BODY, bold=True)
    add_text(slide4, 640080, y+250000, 5000000, 300000, desc, 12, C_MUTED)

for i, (name, desc) in enumerate(industries_right):
    y = row_ys[i]
    add_text(slide4, 6200000, y, 5000000, 300000, name, 16, C_BODY, bold=True)
    add_text(slide4, 6200000, y+250000, 5000000, 300000, desc, 12, C_MUTED)

for dy in [2400000, 3500000]:
    divider(slide4, 640080, dy, 5000000)
    divider(slide4, 6200000, dy, 5000000)

# ── SLIDE 5: Project Planning ────────────────────────────────────────────────
feature_slide(
    "PROJECT PLANNING", "05",
    "Structured Project Planning from Day One",
    [
        "• Hierarchical WBS with unlimited nesting",
        "• Baseline vs. actual tracking",
        "• Milestone and dependency management",
        "• Real-time schedule updates",
    ],
    "Teams complete projects 30% faster with clear ownership and structured planning built in from day one."
)

# ── SLIDE 6: Personal Productivity ──────────────────────────────────────────
feature_slide(
    "PERSONAL PRODUCTIVITY", "06",
    "Every Team Member Knows What to Do",
    [
        "• Personal task dashboard with priorities",
        "• Daily and weekly work planning",
        "• Progress tracking by individual",
        "• Seamless manager visibility",
    ],
    "Individual contributors spend less time in meetings and more time executing — with full clarity on priorities."
)

# ── SLIDE 7: Management Dashboard ───────────────────────────────────────────
feature_slide(
    "MANAGEMENT DASHBOARD", "07",
    "Real-Time View Across All Projects",
    [
        "• Portfolio-level project health at a glance",
        "• Resource utilization heatmaps",
        "• Budget vs. actual tracking",
        "• One-click drill-down to task level",
    ],
    "Managers make faster, better decisions with live data instead of waiting for status reports."
)

# ── SLIDE 8: AI & Intelligence ───────────────────────────────────────────────
slide8 = add_slide()
top_bar(slide8)
section_header(slide8, "AI & INTELLIGENCE", "08")
add_text(slide8, 640080, 700000, 10000000, 400000, "Intelligence Built Into Every Workflow", 28, C_BODY, bold=True)

ai_rows = [
    (1600000, "Smart Scheduling", "AI-suggested timelines based on team capacity and past performance"),
    (2700000, "Risk Detection", "Automatically flags at-risk tasks before they cause delays"),
    (3800000, "Workload Balancing", "Recommends reallocation when team members are over or under-utilized"),
]
for y, title, desc in ai_rows:
    add_oval(slide8, 640080, y, 457200, 457200, C_SECNUM)
    add_text(slide8, 1200000, y, 10500000, 300000, title, 16, C_BODY, bold=True)
    add_text(slide8, 1200000, y+300000, 10500000, 300000, desc, 13, C_MUTED)

for dy in [2500000, 3600000]:
    divider(slide8, 640080, dy)

# ── SLIDE 9: Approvals & Governance ─────────────────────────────────────────
feature_slide(
    "APPROVALS & GOVERNANCE", "09",
    "Compliant, Auditable Approval Workflows",
    [
        "• Multi-level approval chains",
        "• Auto-escalation on SLA breach",
        "• Full audit trail for every action",
        "• Role-based access and visibility",
    ],
    "Governance without the bottleneck — structured approvals that move at the speed of business."
)

# ── SLIDE 10: Why Different ──────────────────────────────────────────────────
slide10 = add_slide()
top_bar(slide10)
section_header(slide10, "WHY DIFFERENT", "10")
add_text(slide10, 640080, 700000, 10000000, 400000, "What Sets Us Apart", 28, C_BODY, bold=True)

diffs = [
    (1500000, "Single Platform", "One tool for planning, execution, tracking, and reporting — no integrations needed"),
    (2400000, "Role-Aware", "Every user sees exactly what they need: tasks, dashboards, or reports"),
    (3300000, "Industry-Agnostic", "Pre-configured for IT, construction, manufacturing, and more"),
    (4200000, "AI-Native", "Intelligence built into workflows, not bolted on"),
]
for y, label, desc in diffs:
    add_text(slide10, 640080, y, 2500000, 400000, label, 16, C_BLUE, bold=True)
    add_text(slide10, 3200000, y, 8350000, 400000, desc, 14, C_BODY)

for dy in [2200000, 3100000, 4000000]:
    divider(slide10, 640080, dy)

# ── SLIDE 11: Business Outcomes ──────────────────────────────────────────────
slide11 = add_slide()
top_bar(slide11)
section_header(slide11, "BUSINESS OUTCOMES", "11")
add_text(slide11, 640080, 700000, 10000000, 400000, "Measurable Impact at Every Level", 28, C_BODY, bold=True)

add_rect(slide11, 6000000, 1300000, 12700, 4000000, C_DIVIDER)

add_text(slide11, 640080, 1300000, 5100000, 200000, "OPERATIONAL", 12, C_BLUE, bold=True)
left_bullets = [
    "• 40% reduction in project delays",
    "• 3x faster planning cycles",
    "• 60% less time on status reporting",
    "• Full resource visibility across teams",
]
tf11l = add_textbox(slide11, 640080, 1600000, 5100000, 3000000)
first = True
for b in left_bullets:
    if first:
        para = tf11l.paragraphs[0]
        first = False
    else:
        para = tf11l.add_paragraph()
    set_para(para, b, 13, C_BODY, space_before=6, space_after=6)

add_text(slide11, 6200000, 1300000, 5500000, 200000, "STRATEGIC", 12, C_BLUE, bold=True)
right_bullets = [
    "• Faster time-to-market",
    "• Higher project success rate",
    "• Reduced operational risk",
    "• Scalable across 10 to 10,000 users",
]
tf11r = add_textbox(slide11, 6200000, 1600000, 5500000, 3000000)
first = True
for b in right_bullets:
    if first:
        para = tf11r.paragraphs[0]
        first = False
    else:
        para = tf11r.add_paragraph()
    set_para(para, b, 13, C_BODY, space_before=6, space_after=6)

# ── SLIDE 12: Vision ─────────────────────────────────────────────────────────
slide12 = add_slide()
add_rect(slide12, 0, 0, 12192000, 6858000, C_LIGHT)
top_bar(slide12)
add_text(slide12, 1000000, 1800000, 10192000, 2000000,
    '"A world where every team has the clarity, tools, and intelligence to deliver — on time, every time."',
    28, C_BODY, italic=True, align=PP_ALIGN.CENTER)
add_rect(slide12, 4500000, 3900000, 3192000, 12700, C_BLUE)
add_text(slide12, 1000000, 4100000, 10192000, 300000,
    "— Inventia Product Vision", 14, C_MUTED, align=PP_ALIGN.CENTER)

# ── SLIDE 13: Thank You ──────────────────────────────────────────────────────
slide13 = add_slide()
add_rect(slide13, 0, 0, 76200, 6858000, C_BLUE)
add_text(slide13, 609600, 2000000, 10000000, 600000, "Let's Build Better Together", 40, C_BODY, bold=True)
add_rect(slide13, 609600, 2800000, 5000000, 12700, C_BLUE)
add_text(slide13, 609600, 3100000, 8000000, 300000, "www.inventia.com | contact@inventia.com", 16, C_BLUE)
add_text(slide13, 609600, 3600000, 3000000, 300000, "Inventia", 14, C_MUTED)

prs.save(OUTPUT_PATH)
print(f"Done: {len(prs.slides)} slides")
